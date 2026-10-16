from __future__ import annotations

import csv
import io
import json
from datetime import datetime, timezone

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from rapidfuzz import fuzz
from sqlalchemy import select
from sqlalchemy.orm import Session

from backend.app.adapters.ashyq import AshyqAdapter
from backend.app.ai.extractor import extract_payment_offer
from backend.app.database import get_db
from backend.app.domain.scoring import calculate_score, has_readiness_mismatch, readiness_gap
from backend.app.models import AuditLog, Evidence, HousingObject, Publication, Signal
from backend.app.schemas import (
    EvidenceCreate,
    EvidenceReview,
    LegalBasisUpdate,
    ObjectCreate,
    PublicationCreate,
    PublicationLinkUpdate,
    SignalCreate,
    SignalTransition,
)
from backend.seed.seed import load_demo_data

router = APIRouter(prefix="/api")
ALLOWED_TRANSITIONS = {
    "draft": {"in_review"},
    "in_review": {"confirmed", "returned"},
    "returned": {"in_review"},
    "confirmed": set(),
}
SOURCE_LINKS = [
    {
        "id": "homeportal",
        "name": "Жилищный портал КЖК",
        "url": "https://homeportal.kz",
        "status": "manual_review",
    },
    {"id": "elicense", "name": "E-License", "url": "https://elicense.kz", "status": "manual_review"},
    {"id": "qazreestr", "name": "Qazreestr", "url": "https://qazreestr.kz", "status": "manual_review"},
    {
        "id": "ashyq",
        "name": "Ashyq Data · Казреестр",
        "url": "https://ashyq.data.gov.kz/dataset/magda-ds-b80110e9-218a-42ba-b8c8-01fa5e4c82c8",
        "status": "api_metadata",
    },
    {
        "id": "equrylys",
        "name": "eQurylys",
        "url": "https://reestr.equrylys.kz/organizations",
        "status": "manual_review",
    },
    {"id": "egov", "name": "eGov · юридические лица", "url": "https://egov.kz", "status": "manual_review"},
    {
        "id": "kgd",
        "name": "КГД · налоговая задолженность",
        "url": "https://kgd.gov.kz",
        "status": "manual_review",
    },
    {
        "id": "rsp",
        "name": "Реестр субъектов предпринимательства",
        "url": "https://rsp.gov.kz",
        "status": "manual_review",
    },
    {
        "id": "cadastre",
        "name": "Публичная кадастровая карта",
        "url": "https://map.gov4c.kz/egkn/",
        "status": "manual_review",
    },
    {
        "id": "procurement",
        "name": "РНУ госзакупок",
        "url": "https://procurement.gov.kz/ru/registry/rnu",
        "status": "manual_review",
    },
]


def _log(
    db: Session,
    action: str,
    housing_object_id: int | None = None,
    entity_type: str = "object",
    entity_id: int | None = None,
    actor: str = "Система",
    details: dict | None = None,
):
    db.add(
        AuditLog(
            housing_object_id=housing_object_id,
            entity_type=entity_type,
            entity_id=entity_id,
            action=action,
            actor=actor,
            details=details or {},
        )
    )


def _iso(value: datetime | None) -> str | None:
    if value is None:
        return None
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    return value.astimezone(timezone.utc).isoformat()


def _score(obj: HousingObject) -> dict:
    return calculate_score(
        [
            {
                "id": e.id,
                "factor_code": e.factor_code,
                "status": e.status,
                "quote": e.quote,
                "source": e.source,
                "source_url": e.source_url,
                "captured_at": e.captured_at.isoformat(),
                "review_reason": e.rationale,
            }
            for e in obj.evidence
        ],
        obj.legal_basis_status,
    )


def _object_json(obj: HousingObject) -> dict:
    return {
        "id": obj.id,
        "name": obj.name,
        "developer": obj.developer,
        "bin": obj.bin,
        "address": obj.address,
        "region": obj.region,
        "status": obj.status,
        "readiness": obj.readiness,
        "declared_readiness": obj.declared_readiness,
        "legal_basis_status": obj.legal_basis_status,
        "legal_basis_type": obj.legal_basis_type,
        "legal_source_url": obj.legal_source_url,
        "latitude": obj.latitude,
        "longitude": obj.longitude,
        "is_demo": obj.is_demo,
        "score": _score(obj),
        "created_at": _iso(obj.created_at),
    }


def _evidence_json(item: Evidence) -> dict:
    return {
        "id": item.id,
        "factor_code": item.factor_code,
        "status": item.status,
        "quote": item.quote,
        "source": item.source,
        "source_url": item.source_url,
        "rationale": item.rationale,
        "captured_at": _iso(item.captured_at),
        "reviewed_by": item.reviewed_by,
        "reviewed_at": _iso(item.reviewed_at),
        "publication_id": item.publication_id,
    }


@router.get("/health")
def health():
    return {"status": "ok", "application": "SAQ QURYLYS"}


@router.get("/objects")
def objects(db: Session = Depends(get_db)):
    return [_object_json(o) for o in db.scalars(select(HousingObject).order_by(HousingObject.id)).all()]


@router.post("/objects", status_code=201)
def create_object(payload: ObjectCreate, db: Session = Depends(get_db)):
    obj = HousingObject(**payload.model_dump())
    db.add(obj)
    db.flush()
    if has_readiness_mismatch(obj.declared_readiness, obj.readiness):
        gap = readiness_gap(obj.declared_readiness, obj.readiness)
        db.add(
            Evidence(
                housing_object_id=obj.id,
                factor_code="R5",
                status="ai_hypothesis",
                source="analyst_readiness_entry",
                rationale=(
                    f"Заявленная готовность {obj.declared_readiness:g}%; "
                    f"введённая подтверждённая готовность {obj.readiness:g}%; "
                    f"разница {gap:g} п.п. Проверьте даты и первичные источники."
                ),
            )
        )
    _log(db, "created", obj.id, details={"name": obj.name})
    db.commit()
    db.refresh(obj)
    return _object_json(obj)


@router.get("/objects/{object_id}")
def object_detail(object_id: int, db: Session = Depends(get_db)):
    obj = db.get(HousingObject, object_id)
    if not obj:
        raise HTTPException(404, "Объект не найден")
    return {
        **_object_json(obj),
        "evidence": [_evidence_json(e) for e in obj.evidence],
        "publications": [
            {
                "id": p.id,
                "source": p.source,
                "source_url": p.source_url,
                "text": p.text,
                "published_at": _iso(p.published_at),
                "extraction": p.extraction,
            }
            for p in obj.publications
        ],
    }


@router.patch("/objects/{object_id}/legal-basis")
def update_legal_basis(object_id: int, payload: LegalBasisUpdate, db: Session = Depends(get_db)):
    obj = db.get(HousingObject, object_id)
    if not obj:
        raise HTTPException(404, "Объект не найден")
    previous = obj.legal_basis_status
    obj.legal_basis_status = payload.status
    obj.legal_basis_type = payload.basis_type if payload.status == "confirmed" else None
    obj.legal_source_url = payload.source_url
    if payload.status == "absent":
        db.add(
            Evidence(
                housing_object_id=obj.id,
                factor_code="R1",
                status="fact",
                source="manual_registry_check",
                source_url=payload.source_url,
                rationale=payload.note,
                reviewed_by=payload.analyst,
                reviewed_at=datetime.now(timezone.utc),
            )
        )
    elif payload.status in {"confirmed", "not_checked"}:
        for item in obj.evidence:
            if item.factor_code == "R1" and item.status == "fact":
                item.status = "rejected"
                item.reviewed_by = payload.analyst
                item.reviewed_at = datetime.now(timezone.utc)
    _log(
        db,
        "legal_basis_updated",
        obj.id,
        actor=payload.analyst,
        details={
            "from": previous,
            "to": payload.status,
            "basis_type": payload.basis_type,
            "source_url": payload.source_url,
            "note": payload.note,
        },
    )
    db.commit()
    db.refresh(obj)
    return _object_json(obj)


@router.get("/objects/{object_id}/audit")
def audit(object_id: int, db: Session = Depends(get_db)):
    if not db.get(HousingObject, object_id):
        raise HTTPException(404, "Объект не найден")
    rows = db.scalars(
        select(AuditLog).where(AuditLog.housing_object_id == object_id).order_by(AuditLog.created_at.desc())
    ).all()
    return [
        {
            "id": row.id,
            "action": row.action,
            "entity_type": row.entity_type,
            "entity_id": row.entity_id,
            "actor": row.actor,
            "details": row.details,
            "created_at": _iso(row.created_at),
        }
        for row in rows
    ]


@router.get("/matches")
def matches(q: str, db: Session = Depends(get_db)):
    candidates = []
    for obj in db.scalars(select(HousingObject)).all():
        name_score = fuzz.WRatio(q, obj.name)
        developer_score = fuzz.WRatio(q, obj.developer) if obj.developer else 0
        confidence = max(name_score, developer_score)
        if confidence >= 35:
            candidates.append(
                {
                    "object_id": obj.id,
                    "name": obj.name,
                    "developer": obj.developer,
                    "confidence": round(confidence / 100, 3),
                    "score": _score(obj),
                }
            )
    return sorted(candidates, key=lambda x: x["confidence"], reverse=True)[:10]


@router.post("/publications", status_code=201)
async def create_publication(payload: PublicationCreate, db: Session = Depends(get_db)):
    if payload.housing_object_id and not db.get(HousingObject, payload.housing_object_id):
        raise HTTPException(404, "Объект не найден")
    parsed = await extract_payment_offer(payload.text)
    obj = db.get(HousingObject, payload.housing_object_id) if payload.housing_object_id else None
    publication = Publication(
        housing_object_id=payload.housing_object_id,
        source=payload.source,
        source_url=payload.source_url,
        text=payload.text,
        published_at=payload.published_at,
        extraction=parsed.model_dump(),
    )
    db.add(publication)
    db.flush()
    evidence = []
    if obj and parsed.has_payment_offer:
        item = Evidence(
            housing_object_id=obj.id,
            publication_id=publication.id,
            factor_code="R2",
            status="ai_hypothesis",
            quote=parsed.quote,
            source=payload.source,
            source_url=payload.source_url,
            rationale=f"Гипотеза {parsed.mode}, уверенность {parsed.confidence:.0%}: {parsed.reasoning}",
        )
        db.add(item)
        db.flush()
        evidence.append(item)
    if obj and obj.status in {"construction", "under_construction", "строится"}:
        item = Evidence(
            housing_object_id=obj.id,
            publication_id=publication.id,
            factor_code="R3",
            status="ai_hypothesis",
            quote=payload.text[:300],
            source=payload.source,
            source_url=payload.source_url,
            rationale="Реклама связана с объектом; проверьте, что объект ещё строится.",
        )
        db.add(item)
        db.flush()
        evidence.append(item)
    _log(
        db,
        "publication_added",
        obj.id if obj else None,
        entity_type="publication",
        entity_id=publication.id,
        details={
            "source": payload.source,
            "has_payment_offer": parsed.has_payment_offer,
            "mode": parsed.mode,
        },
    )
    db.commit()
    return {
        "publication_id": publication.id,
        "housing_object_id": publication.housing_object_id,
        "extraction": parsed.model_dump(),
        "evidence": [_evidence_json(e) for e in evidence],
        "matches": matches(payload.text[:200], db),
    }


@router.patch("/publications/{publication_id}/object")
def link_publication(publication_id: int, payload: PublicationLinkUpdate, db: Session = Depends(get_db)):
    publication = db.get(Publication, publication_id)
    obj = db.get(HousingObject, payload.housing_object_id)
    if not publication:
        raise HTTPException(404, "Публикация не найдена")
    if not obj:
        raise HTTPException(404, "Объект не найден")
    publication.housing_object_id = obj.id
    extraction = publication.extraction or {}
    added = []
    if extraction.get("has_payment_offer") and not db.scalar(
        select(Evidence.id).where(Evidence.publication_id == publication.id, Evidence.factor_code == "R2")
    ):
        item = Evidence(
            housing_object_id=obj.id,
            publication_id=publication.id,
            factor_code="R2",
            status="ai_hypothesis",
            quote=extraction.get("quote"),
            source=publication.source,
            source_url=publication.source_url,
            rationale=(
                f"Гипотеза {extraction.get('mode', 'rules')}, "
                f"уверенность {float(extraction.get('confidence', 0)):.0%}: "
                f"{extraction.get('reasoning', '')}"
            ),
        )
        db.add(item)
        db.flush()
        added.append(item)
    if obj.status in {"construction", "under_construction", "строится"} and not db.scalar(
        select(Evidence.id).where(Evidence.publication_id == publication.id, Evidence.factor_code == "R3")
    ):
        item = Evidence(
            housing_object_id=obj.id,
            publication_id=publication.id,
            factor_code="R3",
            status="ai_hypothesis",
            quote=publication.text[:300],
            source=publication.source,
            source_url=publication.source_url,
            rationale="Реклама связана с объектом; проверьте, что объект ещё строится.",
        )
        db.add(item)
        db.flush()
        added.append(item)
    _log(
        db,
        "publication_linked",
        obj.id,
        "publication",
        publication.id,
        details={"object_name": obj.name},
    )
    db.commit()
    return {
        "publication_id": publication.id,
        "housing_object_id": obj.id,
        "evidence": [_evidence_json(item) for item in added],
    }


@router.post("/objects/{object_id}/evidence", status_code=201)
def add_evidence(object_id: int, payload: EvidenceCreate, db: Session = Depends(get_db)):
    obj = db.get(HousingObject, object_id)
    if not obj:
        raise HTTPException(404, "Объект не найден")
    if payload.factor_code == "R1":
        raise HTTPException(422, "R1 фиксируется только через проверку правового основания с источником")
    item = Evidence(
        housing_object_id=obj.id,
        factor_code=payload.factor_code,
        status="ai_hypothesis",
        quote=payload.quote,
        source=payload.source,
        source_url=payload.source_url,
        rationale=payload.rationale,
    )
    db.add(item)
    db.flush()
    _log(
        db,
        "evidence_hypothesis_added",
        obj.id,
        "evidence",
        item.id,
        details={"factor_code": item.factor_code},
    )
    db.commit()
    db.refresh(item)
    return _evidence_json(item)


@router.patch("/evidence/{evidence_id}/review")
def review_evidence(evidence_id: int, payload: EvidenceReview, db: Session = Depends(get_db)):
    item = db.get(Evidence, evidence_id)
    if not item:
        raise HTTPException(404, "Фактор не найден")
    if payload.source_url:
        item.source_url = payload.source_url
    if payload.status == "fact" and not item.source_url:
        raise HTTPException(422, "Для подтверждения фактора сохраните ссылку на источник")
    if payload.status == "fact" and item.factor_code in {"R2", "R3"} and not item.quote:
        raise HTTPException(422, "Для подтверждения фактора рекламы нужна точная цитата")
    item.status = payload.status
    item.reviewed_by = payload.reviewer
    item.reviewed_at = datetime.now(timezone.utc)
    if payload.comment:
        item.rationale = f"{item.rationale}\nКомментарий аналитика: {payload.comment}".strip()
    _log(
        db,
        "evidence_reviewed",
        item.housing_object_id,
        "evidence",
        item.id,
        payload.reviewer,
        {"status": payload.status, "comment": payload.comment},
    )
    db.commit()
    db.refresh(item)
    return {"evidence": _evidence_json(item), "score": _score(item.housing_object)}


@router.get("/sources")
async def sources():
    result = []
    for source in SOURCE_LINKS:
        if source["id"] == "ashyq":
            result.append(await AshyqAdapter().fetch())
        else:
            result.append(
                {
                    **source,
                    "message": "Источник открыт для ручной проверки; автоматический доступ не настроен.",
                }
            )
    return result


@router.post("/import/ashyq", status_code=201)
async def import_ashyq(file: UploadFile = File(...), db: Session = Depends(get_db)):
    if file.filename is None or not file.filename.lower().endswith((".csv", ".xlsx", ".xlsm")):
        raise HTTPException(415, "Загрузите CSV или XLSX из набора Ashyq Data")
    raw = await file.read(5_000_001)
    if len(raw) > 5_000_000:
        raise HTTPException(413, "Файл больше 5 МБ")
    try:
        if file.filename.lower().endswith(".csv"):
            text = raw.decode("utf-8-sig")
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
            rows = list(csv.DictReader(io.StringIO(text), dialect=dialect))
        else:
            from openpyxl import load_workbook

            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            sheet = workbook.active
            values = sheet.iter_rows(values_only=True)
            headers = [str(v or "").strip() for v in next(values)]
            rows = [dict(zip(headers, row)) for row in values]
    except (UnicodeDecodeError, csv.Error, StopIteration, ValueError) as error:
        raise HTTPException(422, f"Не удалось разобрать файл: {type(error).__name__}") from error
    if len(rows) > 25000:
        raise HTTPException(413, "В файле больше 25 000 строк")
    imported = 0
    duplicates = 0
    for row in rows:
        # Ashyq's published schema uses object_id/object_address, not object_name.
        # Keep legacy aliases for manually prepared spreadsheets as well.
        object_id = str(row.get("object_id") or row.get("id") or "").strip()
        address = str(
            row.get("object_address")
            or row.get("address_name")
            or row.get("Адрес")
            or ""
        ).strip()
        location = str(row.get("object_location") or "").strip()
        name = str(row.get("object_name") or row.get("Название объекта") or address or location).strip()
        if not name and object_id:
            name = f"Объект {object_id}"
        if not name:
            continue
        bin_value = str(row.get("bin") or row.get("БИН") or "").strip()
        bin_value = bin_value if len(bin_value) == 12 and bin_value.isdigit() else None
        developer = str(
            row.get("builder_name") or row.get("Застройщик") or row.get("authorized_org_name") or ""
        ).strip()
        exists = db.scalar(
            select(HousingObject).where(
                HousingObject.name == name,
                HousingObject.developer == developer,
                HousingObject.address == address,
            )
        )
        if exists:
            duplicates += 1
            continue
        obj = HousingObject(
            name=name,
            developer=developer,
            bin=bin_value,
            address=address,
            region="",
            status="unknown",
            legal_basis_status="not_checked",
            is_demo=False,
        )
        db.add(obj)
        db.flush()
        _log(
            db,
            "imported_ashyq",
            obj.id,
            actor="Импорт Ashyq",
            details={
                "source_id": object_id or None,
                "permission_number": str(row.get("permission_number") or ""),
                "date_start": str(row.get("date_start") or ""),
                "date_end": str(row.get("date_end") or ""),
                "time_modif": str(row.get("time_modif") or ""),
            },
        )
        imported += 1
    db.commit()
    return {
        "uploaded_rows": len(rows),
        "imported_objects": imported,
        "duplicates_skipped": duplicates,
        "message": "Импортированы записи источника. Импорт не означает наличия риска и не создаёт R1.",
    }


@router.get("/scenarios")
def scenarios():
    path = __import__("pathlib").Path(__file__).resolve().parents[2] / "seed" / "scenarios.json"
    return json.loads(path.read_text(encoding="utf-8"))


@router.post("/demo-seed")
def seed_demo_data():
    counts = load_demo_data()
    return {
        "message": "Загружены вымышленные контрольные данные с маркировкой ДЕМО.",
        **counts,
    }


@router.get("/signals")
def signals(db: Session = Depends(get_db)):
    rows = db.scalars(select(Signal).order_by(Signal.updated_at.desc())).all()
    return [
        {
            "id": s.id,
            "housing_object_id": s.housing_object_id,
            "object_name": s.housing_object.name,
            "status": s.status,
            "author": s.author,
            "comments": s.comments,
            "score_snapshot": s.score_snapshot,
            "updated_at": _iso(s.updated_at),
        }
        for s in rows
    ]


@router.post("/signals", status_code=201)
def create_signal(payload: SignalCreate, db: Session = Depends(get_db)):
    obj = db.get(HousingObject, payload.housing_object_id)
    if not obj:
        raise HTTPException(404, "Объект не найден")
    signal = Signal(
        housing_object_id=obj.id,
        author=payload.author,
        comments=(
            [
                {
                    "author": payload.author,
                    "text": payload.comment,
                    "created_at": datetime.now(timezone.utc).isoformat(),
                }
            ]
            if payload.comment
            else []
        ),
        score_snapshot=_score(obj),
    )
    db.add(signal)
    db.flush()
    _log(
        db,
        "signal_created",
        obj.id,
        "signal",
        signal.id,
        payload.author,
        {"score": signal.score_snapshot.get("score")},
    )
    db.commit()
    return {"id": signal.id, "status": signal.status, "score_snapshot": signal.score_snapshot}


@router.patch("/signals/{signal_id}/status")
def transition_signal(signal_id: int, payload: SignalTransition, db: Session = Depends(get_db)):
    signal = db.get(Signal, signal_id)
    if not signal:
        raise HTTPException(404, "Сигнальная карточка не найдена")
    if payload.status not in ALLOWED_TRANSITIONS[signal.status]:
        raise HTTPException(409, f"Переход {signal.status} → {payload.status} запрещён")
    if payload.status == "confirmed" and not signal.score_snapshot.get("factors"):
        raise HTTPException(422, "Карточка без подтверждённых факторов не может быть утверждена")
    signal.status = payload.status
    signal.comments = [
        *signal.comments,
        {
            "author": payload.author,
            "text": payload.comment,
            "created_at": datetime.now(timezone.utc).isoformat(),
        },
    ]
    _log(
        db,
        "signal_status_changed",
        signal.housing_object_id,
        "signal",
        signal.id,
        payload.author,
        {"status": payload.status, "comment": payload.comment},
    )
    db.commit()
    return {"id": signal.id, "status": signal.status, "comments": signal.comments}
