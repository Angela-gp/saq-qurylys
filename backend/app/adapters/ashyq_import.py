from __future__ import annotations

import csv
import hashlib
import io
import json
import re
from pathlib import Path
from typing import Any

CONFIG_PATH = Path(__file__).with_name("ashyq_columns.json")
DATASET_NAME = "Разрешения на привлечение денег дольщиков (Ashyq Data)"
DATASET_URL = "https://ashyq.data.gov.kz/dataset/magda-ds-b80110e9-218a-42ba-b8c8-01fa5e4c82c8"


def normalize_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9а-яё]+", "", str(value or "").casefold())


def parse_rows(raw: bytes, filename: str) -> list[dict[str, Any]]:
    if filename.lower().endswith(".csv"):
        text = raw.decode("utf-8-sig")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        return [dict(row) for row in csv.DictReader(io.StringIO(text), dialect=dialect)]
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    try:
        sheet = workbook.active
        values = sheet.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(values)]
        return [dict(zip(headers, row)) for row in values if any(value is not None for value in row)]
    finally:
        workbook.close()


def map_row(row: dict[str, Any]) -> dict[str, Any]:
    config = json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
    normalized = {normalize_header(key): value for key, value in row.items()}
    result: dict[str, Any] = {}
    for field, aliases in config.items():
        for alias in aliases:
            value = normalized.get(normalize_header(alias))
            if value is not None and str(value).strip():
                result[field] = str(value).strip()
                break
    result["name"] = result.get("name") or result.get("address") or result.get("location")
    result["developer"] = result.get("developer", "")
    result["address"] = result.get("address") or result.get("location", "")
    result["bin"] = result.get("bin", "")
    result["bin"] = result["bin"] if len(result["bin"]) == 12 and result["bin"].isdigit() else None
    return result


def record_key(mapped: dict[str, Any]) -> str:
    source_record_id = mapped.get("source_record_id")
    if source_record_id:
        basis = f"ashyq:{source_record_id}"
    else:
        basis = "ashyq:" + "|".join(
            str(mapped.get(key, "")).casefold().strip()
            for key in ("name", "developer", "address", "permission_number")
        )
    return hashlib.sha256(basis.encode("utf-8")).hexdigest()


def looks_like_person_name(value: str | None) -> bool:
    text = (value or "").strip()
    if not text:
        return False
    lowered = text.casefold()
    if lowered.startswith(("ип ", "ип\u00a0", "жк ", "жеке кәсіпкер")):
        return True
    parts = re.findall(r"[^\W\d_]+", text, flags=re.UNICODE)
    if not 2 <= len(parts) <= 3 or len(" ".join(parts)) != len(text):
        return False
    surname_endings = ("ов", "ев", "ин", "ова", "ева", "ина", "ұлы", "қызы", "ович", "евна", "овна")
    return any(parts[-1].casefold().endswith(ending) for ending in surname_endings)
